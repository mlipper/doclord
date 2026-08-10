"""
Extracts contact information from Excel request forms.

Distinguishes between v1 and v2 signup forms because contact information
exists in different cell locations:

form v1:
  B16:E20
  Bxx - name
  Cxx - email
  Dxx - phone
  Exx - role
    
form v2:
  C25:F29
  Cxx - name
  Dxx - email
  Exx - phone
  Fxx - role
"""
import logging

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import ClassVar

import pandas as pd
from openpyxl.reader.excel import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


_log = logging.getLogger(__name__)


APP_COLUMNS_V1 = ["B", "C", "D", "E"]
APP_COLUMNS_V2 = ["C", "D", "E", "F"]
APP_DISPLAY_NAME_V1 = "C8"
APP_DISPLAY_NAME_V2 = "C34"
APP_ROWS_V1 = [16, 17, 18, 19, 20]
APP_ROWS_V2 = [25, 26, 27, 28, 29]
FORM_V1 = "Geoclient API - v1"
FORM_V2 = "Geoclient API"
INPUT_DIR = "files"
OUTPUT_DIR = "out"
ROLE_SPONSOR = "Sponsor"
ROLE_SUPPORT = "Support"
ROLE_TEAM_DEFAULT = "Team"
SHEET_NAME = "Request Form"
VERSION_CELL = "B3"


@dataclass
class Contact:
    """Generic record-like object for contact information."""
    email: str
    name: str = None
    phone: str = None
    role: str = None

    def has_info(self) -> bool:
        return self.email is not None or self.name is not None or self.phone is not None

@dataclass
class Application:
    """Container for an application's contact information."""
    display_name: str
    name: str
    team: list[Contact] = field(default_factory=list)
    COLUMNS: ClassVar[list[str]] = ["display_name", "application", "email", "name", "phone", "role"] 

    def asrows(self):
        rows = []
        for c in self.team:
            rows.append(
                [
                self.sstrip(self.display_name),
                self.sstrip(self.name),
                self.sstrip(c.email),
                self.sstrip(c.name),
                self.sstrip(f"{c.phone}"),
                self.sstrip(c.role)
                ]
                )
        return rows

    def sstrip(self, string: str):
        """Safe str.strip() in case str is None"""
        return string.strip() if string is not None else None


def _cval(ws: Worksheet, cell: str, default_value=None):
    if ws is None:
        return default_value
    c = ws[cell]
    if c is None:
        return default_value
    v = c.value
    if v is None:
        return default_value
    return v

def _get_display_name(ws: Worksheet, form: str):
    """
    Gets the display name from the Worksheet using the cell address.
    """
    if form == FORM_V2:
        return ws[APP_DISPLAY_NAME_V2].value

    return ws[APP_DISPLAY_NAME_V1].value

def _get_group(ws: Worksheet):
    """
    Gets the distribution list contact info.
    Only applies to the v2 form.

    C23 - group email
    D23 - group name
    E23 - group phone

    Parameters
    ----------
    ws: Worksheet

    Returns
    -------
    Contact
    """
    contact = Contact(_cval(ws, "C23"), _cval(ws, "D23"), _cval(ws, "E23"), ROLE_SUPPORT)
    if not contact.has_info():
        return None
    return contact

def _get_sponsor(ws: Worksheet):
    """
    Gets the project sponsor contact info.
    Only applies to the v2 form.

    C18 - name
    D18 - email
    E18 - phone
    F18 - title

    Parameters
    ----------
    ws: Worksheet

    Returns
    -------
    Contact
    """
    contact = Contact(_cval(ws, "D18"), _cval(ws, "C18"), _cval(ws, "E18"), _cval(ws, "F18", ROLE_SPONSOR))
    if not contact.has_info():
        return None
    return contact

def _get_team(ws: Worksheet, form: str = FORM_V2):
    """
    For the v2 form, extracts the application team members contact info.
    For the v1 form, extracts all contact info.

    Parameters
    ----------
    ws: Worksheet containing the populated form. 
    form: Form version 1 or 2 (default).

    Returns
    -------
    list of Contact
    """
    rows = APP_ROWS_V2
    columns = APP_COLUMNS_V2

    if form == FORM_V1:
        rows = APP_ROWS_V1
        columns = APP_COLUMNS_V1

    team_members = []
    for row in rows:
        email = None
        name = None
        phone = None
        role = None
        for col in columns:
            match columns.index(col):
                case 0:
                    name = _cval(ws,f"{col}{row}")
                case 1:
                    email = _cval(ws,f"{col}{row}")
                case 2:
                    phone = _cval(ws,f"{col}{row}")
                case 3:
                    role = _cval(ws,f"{col}{row}", ROLE_TEAM_DEFAULT)
                case _:
                    raise ValueError(f"Unexpected column {col}.")
        contact = Contact(email, name, phone, role)
        if contact.has_info():
            team_members.append(contact)
    return team_members

def timestamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S")

def _write_report(apps: list[Application], io: Path):
    data = []
    for app in apps:
        #for contact in app.team:
        #    logging.info("%s, %s, %s, %s, %s, %s", app.display_name, app.name, contact.email, contact.name, contact.phone, contact.role)
        for row in app.asrows():
            if len(row) != 6:
                logging.error("%s - %s", app.name, row)
            data.append(row)
    df = pd.DataFrame(data, columns=Application.COLUMNS)
    df.index += 1
    df.to_csv(io)

def slurp(excel_file: Path):
    """
    Extracts contact information from the given Excel file.
    """
    wb = load_workbook(excel_file, True, False, True)
    ws = wb[SHEET_NAME]


    display_name = None
    team = None
    if ws[VERSION_CELL].value == FORM_V2:
        sponsor = _get_sponsor(ws)
        team = _get_team(ws)
        group = _get_group(ws)
        if group is not None and group.has_info():
            team.append(group)
        if sponsor is not None and sponsor.has_info():
            team.append(sponsor)
        display_name = _get_display_name(ws, FORM_V2)
    elif ws[VERSION_CELL].value == FORM_V1:
        team = _get_team(ws, FORM_V1)
        display_name = _get_display_name(ws, FORM_V1)
    else:
        logging.error("Unrecognized form version %s for %s.", ws[VERSION_CELL], excel_file.stem)
        team = []

    return Application(display_name, excel_file.stem, team)

def main():
    logging.basicConfig(level=logging.INFO)

    cwd = Path(__file__).parent
    input_dir = cwd / INPUT_DIR
    xlsx_files = sorted(input_dir.rglob("*.xlsx"))

    apps = []
    for excel_file in xlsx_files:
        app = slurp(excel_file)
        apps.append(app)
    output_dir = cwd / OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / f"geoclient-contacts-{timestamp()}.csv"
    _write_report(apps, output_file)


if __name__ == "__main__":
    print("Hello")
    main()
    print("Goodbye")
